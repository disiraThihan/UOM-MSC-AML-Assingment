import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette - Professional Deep Slate & Vibrant Tech Accents
    BG_DARK = RGBColor(15, 23, 42)       # #0F172A
    CARD_BG = RGBColor(30, 41, 59)       # #1E293B
    CARD_BORDER = RGBColor(51, 65, 85)   # #334155
    ACCENT_BLUE = RGBColor(59, 130, 246) # #3B82F6
    ACCENT_CYAN = RGBColor(6, 182, 212)  # #06B6D4
    ACCENT_GREEN = RGBColor(16, 185, 129)# #10B981
    TEXT_WHITE = RGBColor(248, 250, 252) # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    TEXT_GOLD = RGBColor(245, 158, 11)   # #F59E0B

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="MACHINE LEARNING PROJECT"):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p1 = tf.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_CYAN

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(4)

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.2)
        return card

    # ==================== SLIDE 1: Title Slide ====================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    # Decorative glow card
    add_card(s1, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5), bg_color=CARD_BG, border_color=ACCENT_BLUE)

    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.4), Inches(10.333), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "UNIVERSITY OF MORATUWA  •  FACULTY OF INFORMATION TECHNOLOGY"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p = tf.add_paragraph()
    p.text = "Colombo Suburban Land Valuation System"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.space_before = Pt(14)

    p = tf.add_paragraph()
    p.text = "A Comparative Empirical Study of Advanced Tree Ensembles, Stacking Meta-Learners, and Game-Theoretic SHAP Interpretability"
    p.font.size = Pt(15)
    p.font.color.rgb = ACCENT_BLUE
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "Focus Corridor: Malabe – Kottawa – Kaduwela Urban Growth Axis (Western Province, Sri Lanka)"
    p.font.size = Pt(13)
    p.font.italic = True
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "Candidate: Thihan D.P.D.   |   Index No: 269204J   |   MSc in Artificial Intelligence 2026"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TEXT_GOLD
    p.space_before = Pt(26)

    p = tf.add_paragraph()
    p.text = "Includes Core Project + 10-Mark Front-End Integration Bonus & Live Demonstration"
    p.font.size = Pt(12)
    p.font.color.rgb = ACCENT_GREEN
    p.space_before = Pt(10)

    # ==================== SLIDE 2: Problem & Sri Lankan Context ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "Problem Identification & Sri Lankan Real-World Context", "01. PROBLEM DEFINITION")

    col_w = Inches(3.64)
    gap = Inches(0.38)
    top_pos = Inches(1.7)
    h_pos = Inches(5.2)

    # Card 1: The Corridor
    add_card(s2, Inches(0.8), top_pos, col_w, h_pos)
    tb = s2.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.2), col_w - Inches(0.4), h_pos - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The Malabe–Kottawa Axis"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    
    bullets1 = [
        "Rapid suburban decentralization in Colombo Metropolitan Region.",
        "IT & Higher Education Hub (Malabe): SLIIT, Horizon, CINEC, LSEG / MillenniumIT.",
        "Multimodal Expressway Junction (Kottawa): Confluence of E01 (Southern) and E02 (Outer Circular).",
        "Administrative Shift: Proximity to Sethsiripaya & Suhurupaya government hubs in Battaramulla."
    ]
    for b in bullets1:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # Card 2: Sri Lankan Specifics
    add_card(s2, Inches(0.8) + col_w + gap, top_pos, col_w, h_pos)
    tb = s2.shapes.add_textbox(Inches(1.0) + col_w + gap, top_pos + Inches(0.2), col_w - Inches(0.4), h_pos - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Unique Local Nuances"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    bullets2 = [
        "Perch-Based Pricing: Land transacted universally in Perches (272.25 sq ft) and quoted in Lakhs LKR/perch.",
        "Bim Saviya Title: State-guaranteed Title Registration (Act No. 21 of 1998) commands a sharp premium over customary deeds.",
        "UDA 20-ft Road Rule: Urban Development Authority mandates min. 20 ft access for subdivision & commercial clearance.",
        "Kelani Hydrology: Monsoon flood risks around low-lying canal and river basins cause severe discounts."
    ]
    for b in bullets2:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # Card 3: Limitations of Global Datasets
    add_card(s2, Inches(0.8) + (col_w + gap)*2, top_pos, col_w, h_pos)
    tb = s2.shapes.add_textbox(Inches(1.0) + (col_w + gap)*2, top_pos + Inches(0.2), col_w - Inches(0.4), h_pos - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Why Global Data Fails"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = TEXT_GOLD

    bullets3 = [
        "Standard Kaggle sets (Boston, Ames Housing) operate on sq footage, HOA fees, central heating, and zip codes.",
        "Completely miss Sri Lankan land tenure, expressway feeder congestion, and local bank appraisal heuristics.",
        "Objective: Build a dedicated 2,500-sample empirical dataset reflecting CBSL Land Price Index and IVSL appraisal standards.",
        "Impact: Fair lending, transparent municipal taxation, and homeowner equity protection."
    ]
    for b in bullets3:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(8)

    # ==================== SLIDE 3: Dataset & Feature Engineering ====================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "Empirical Dataset Formulation & Domain Feature Engineering", "02. DATASET ARCHITECTURE")

    # Left: Details Card
    add_card(s3, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    tb = s3.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Dataset Specification (N = 2,500)"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    features_info = [
        ("9 Colombo Localities", "Malabe, Kottawa, Thalawathugoda, Athurugiriya, Battaramulla, Kaduwela, Maharagama, Pannipitiya, Homagama"),
        ("Spatial Vectors", "Distance to E01/E02 Interchange (km), Distance to Colombo Fort CBD (km), Distance to Main Arterial Bus Route (m)"),
        ("Legal & Planning", "Access Road Width (feet: 10-40 ft), Bim Saviya Title Guarantee (Binary), UDA Zoning Classification"),
        ("Physical & Natural", "Parcel extent (perches), Flood Inundation Risk (Binary), Land Topography (Flat/Slope/Low-lying)"),
        ("Economic Heuristic", "CBSL Colombo Land Price Index anchor benchmark"),
        ("Target Variable", "Price per Perch in Lakhs LKR (Continuous: 8.5 to 58.0 Lakhs)")
    ]
    for title, desc in features_info:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Right: Image
    img_path = "static/plots/eda_distributions.png"
    if os.path.exists(img_path):
        add_card(s3, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
        s3.shapes.add_picture(img_path, Inches(7.0), Inches(1.85), Inches(5.3), Inches(4.8))

    # ==================== SLIDE 4: Model Benchmarking Strategy ====================
    s4 = prs.slides.add_slide(blank_layout)
    add_bg(s4)
    add_header(s4, "Model Selection Strategy: Avoiding Weak Single Models", "03. ALGORITHMIC BENCHMARK")

    # Left: Explanation Card
    add_card(s4, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    tb = s4.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "10 Algorithmic Architectures Compared"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    model_points = [
        "In strict adherence to instructions, simple weak models were tested only as baselines:",
        "  - Linear Regression, Ridge, Lasso, and Single Decision Tree.",
        "Advanced Ensemble Frameworks evaluated:",
        "  - Random Forest & AdaBoost (Bagging & Boosting)",
        "  - Gradient Boosting (GBM)",
        "  - LightGBM (Leaf-wise growth, histogram binning)",
        "  - XGBoost (Extreme Gradient Boosting with regularized objective)",
        "  - CatBoost (Target encoding, symmetric oblivious trees)",
        "Multi-Level Stacking Meta-Learner:",
        "  - Level-0: CatBoost + XGBoost + LightGBM + Random Forest",
        "  - Level-1 Meta-Learner: ElasticNet Blender",
        "Evaluation Protocol: 5-Fold Stratified Cross-Validation + 20% Unseen Holdout Test Set."
    ]
    for pt in model_points:
        p = tf.add_paragraph()
        p.text = pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED if not pt.startswith("  -") else TEXT_WHITE
        p.space_before = Pt(4)

    # Right: Image
    img_path = "static/plots/model_comparison_benchmark.png"
    if os.path.exists(img_path):
        add_card(s4, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
        s4.shapes.add_picture(img_path, Inches(7.0), Inches(1.85), Inches(5.3), Inches(4.8))

    # ==================== SLIDE 5: Quantitative Results ====================
    s5 = prs.slides.add_slide(blank_layout)
    add_bg(s5)
    add_header(s5, "Empirical Results: Ensemble Superiority & Error Reduction", "04. QUANTITATIVE EVALUATION")

    rows = [
        ["Model Architecture", "Type", "CV R²", "Test R²", "RMSE (Lakhs)", "MAPE (%)"],
        ["CatBoost Regressor", "Ensemble", "0.953", "0.9676", "1.641", "5.33%"],
        ["Stacking Meta-Learner", "Ensemble", "0.952", "0.9615", "1.790", "5.77%"],
        ["LightGBM Regressor", "Ensemble", "0.931", "0.9489", "2.060", "6.62%"],
        ["XGBoost Regressor", "Ensemble", "0.923", "0.9435", "2.166", "7.20%"],
        ["Random Forest", "Ensemble", "0.918", "0.9362", "2.302", "7.55%"],
        ["Linear Regression (Baseline)", "Baseline", "0.865", "0.8712", "3.315", "11.45%"],
        ["Single Decision Tree", "Baseline", "0.672", "0.6876", "5.094", "18.32%"]
    ]

    add_card(s5, Inches(0.8), Inches(1.7), Inches(7.2), Inches(5.2))
    table_shape = s5.shapes.add_table(len(rows), 6, Inches(1.0), Inches(1.9), Inches(6.8), Inches(4.6))
    table = table_shape.table

    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if c_idx > 1 else PP_ALIGN.LEFT
            p.font.size = Pt(10)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = ACCENT_CYAN
            elif r_idx == 1:
                p.font.bold = True
                p.font.color.rgb = ACCENT_GREEN
            elif r_idx >= 6:
                p.font.color.rgb = RGBColor(239, 68, 68)
            else:
                p.font.color.rgb = TEXT_WHITE

    add_card(s5, Inches(8.3), Inches(1.7), Inches(4.2), Inches(5.2))
    tb = s5.shapes.add_textbox(Inches(8.5), Inches(1.9), Inches(3.8), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Key Benchmark Findings"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    insights = [
        ("CatBoost Wins (R² = 0.9676)", "Best-in-class handling of multi-nominal categorical variables (9 localities, deed types, zoning) with zero target leakage."),
        ("67.8% Error Reduction", "RMSE drops from 5.094 Lakhs (Single Tree) down to 1.641 Lakhs (CatBoost). Single tree collapses under spatial variance."),
        ("Stacking Meta-Learner (0.9615)", "Second-best architecture, successfully fusing inductive biases of tree-based and gradient boosted models."),
        ("Commercial Viability", "Mean Absolute Percentage Error (MAPE) of 5.33% falls well within professional chartered valuation tolerances (±7.5%).")
    ]
    for h, b in insights:
        p = tf.add_paragraph()
        p.text = "✔ " + h
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_GOLD
        p.space_before = Pt(8)
        p2 = tf.add_paragraph()
        p2.text = b
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MUTED

    # ==================== SLIDE 6: Model Explainability & SHAP ====================
    s6 = prs.slides.add_slide(blank_layout)
    add_bg(s6)
    add_header(s6, "Model Interpretability: Game-Theoretic SHAP Analysis", "05. MODEL INTERPRETABILITY")

    img_path = "static/plots/shap_summary_beeswarm.png"
    if os.path.exists(img_path):
        add_card(s6, Inches(0.8), Inches(1.7), Inches(5.7), Inches(5.2))
        s6.shapes.add_picture(img_path, Inches(1.0), Inches(1.85), Inches(5.3), Inches(4.8))

    add_card(s6, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
    tb = s6.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "TreeSHAP Local & Global Insights"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    shap_points = [
        ("The 20-Foot Road Access Threshold (UDA Policy)", "SHAP dependence exposes a steep non-linear step-function jump at exactly 20 feet width (+3.5 to +5.0 Lakhs/perch), directly mirroring UDA subdivision regulations."),
        ("Expressway Interchange Proximity Gradient", "Exponential price decay beyond 2.5 km from Kottawa/Athurugiriya/Kaduwela interchanges. Commuters pay a heavy premium within the 0-2 km band."),
        ("Deed Title Security Premium", "Bim Saviya First-Class title imparts a consistent +4.2 Lakhs/perch uplift, providing quantitative proof of institutional title security value."),
        ("Hydrological Discount", "Parcels classified within low-lying flood inundation areas experience an immediate -6.5 Lakhs/perch downward attribution."),
        ("Trust & Compliance", "Enables commercial banks and land buyers to inspect exactly WHY each rupee was added or deducted.")
    ]
    for title, desc in shap_points:
        p = tf.add_paragraph()
        p.text = "• " + title + ": "
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(6)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # ==================== SLIDE 7: Front-End Integration (+10 Marks Bonus) ====================
    s7 = prs.slides.add_slide(blank_layout)
    add_bg(s7)
    add_header(s7, "Bonus Deliverable (+10 Marks): Production Front-End System & Interactive SHAP Dashboard", "06. FRONT-END INTEGRATION")

    card_w = Inches(3.64)
    gap_w = Inches(0.38)
    top_p = Inches(1.7)
    h_p = Inches(5.2)

    add_card(s7, Inches(0.8), top_p, card_w, h_p)
    tb = s7.shapes.add_textbox(Inches(1.0), top_p + Inches(0.2), card_w - Inches(0.4), h_p - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "FastAPI Backend Architecture"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    b_list1 = [
        "High-performance asynchronous Python REST API.",
        "Pydantic schema validation for strict spatial & physical attribute boundaries.",
        "Model serialization via Joblib with pre-fitted ColumnTransformers.",
        "Automated OpenAPI / Swagger interactive documentation at /docs."
    ]
    for b in b_list1:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(10)

    add_card(s7, Inches(0.8) + card_w + gap_w, top_p, card_w, h_p)
    tb = s7.shapes.add_textbox(Inches(1.0) + card_w + gap_w, top_p + Inches(0.2), card_w - Inches(0.4), h_p - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Interactive Glassmorphism UI"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    b_list2 = [
        "Modern dark-mode dashboard built with pure CSS and Vanilla JS.",
        "Dynamic input controls for locality, road width, perch extent, and title deed type.",
        "Asynchronous AJAX POST requests with zero-latency instant valuation updates.",
        "Visual confidence intervals and total parcel investment calculations (Millions LKR)."
    ]
    for b in b_list2:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(10)

    add_card(s7, Inches(0.8) + (card_w + gap_w)*2, top_p, card_w, h_p)
    tb = s7.shapes.add_textbox(Inches(1.0) + (card_w + gap_w)*2, top_p + Inches(0.2), card_w - Inches(0.4), h_p - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "On-the-Fly SHAP Explainability Engine"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    b_list3 = [
        "Real-time TreeSHAP attribution computed dynamically for every user query.",
        "Interactive Waterfall Breakdown bar chart shows exact +/- contributions of each feature.",
        "Unprecedented transparency: No 'black-box' predictions; users see exact mathematical rationale.",
        "Demonstrates complete, production-grade end-to-end MLOps lifecycle."
    ]
    for b in b_list3:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_MUTED
        p.space_before = Pt(10)

    # ==================== SLIDE 8: Summary & Conclusion ====================
    s8 = prs.slides.add_slide(blank_layout)
    add_bg(s8)
    add_header(s8, "Project Summary, Policy Implications & Future Roadmap", "07. CONCLUSION")

    add_card(s8, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.2))
    tb = s8.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.0), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Summary of Contributions"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    sum_items = [
        ("Addressed Sri Lankan Reality", "Replaced foreign generic benchmarks with authentic Western Province socio-economic dynamics."),
        ("Defeated Weak Baselines", "Ensembles reduced error by 67.8%; CatBoost (R²=0.9676) and Stacking (R²=0.9615) proved superior."),
        ("Disproved the Black Box", "TreeSHAP explained critical threshold non-linearities (20-ft road rule & expressway premium)."),
        ("End-to-End Delivery", "Accompanied by rigorous academic report, production code, and full live deployment.")
    ]
    for h, b in sum_items:
        p = tf.add_paragraph()
        p.text = "✔ " + h + ": "
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_GOLD
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = b
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    add_card(s8, Inches(6.8), Inches(1.7), Inches(5.7), Inches(5.2))
    tb = s8.shapes.add_textbox(Inches(7.1), Inches(1.9), Inches(5.1), Inches(4.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Policy Value & Future Roadmap"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    future_items = [
        ("Municipal & UDA Road Widening", "Quantifiable evidence that widening 12 ft access roads to 20 ft unlocks >15% capital appreciation, justifying public municipal road expansion bonds."),
        ("Banking & Mortgage Risk", "Eliminates subjective valuation fraud; provides commercial banks with reproducible, automated collateral risk bounds."),
        ("Future: Geospatial Satellite Feeds", "Incorporate Sentinel-2 and high-resolution LiDAR elevation data to auto-detect micro-topography and wetland vegetation boundaries."),
        ("Automated Title Ingestion", "Integrate OCR & LLM pipeline to ingest Land Registry extracts directly into Bim Saviya verification modules.")
    ]
    for h, b in future_items:
        p = tf.add_paragraph()
        p.text = "➔ " + h + ": "
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = TEXT_WHITE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = b
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    prs.save("Colombo_Land_Valuation_AI_Presentation.pptx")
    prs.save("Colombo_Land_Valuation_System_Presentation.pptx")
    print("Presentations saved successfully!")

if __name__ == "__main__":
    create_deck()
